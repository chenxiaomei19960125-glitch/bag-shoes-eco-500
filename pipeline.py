"""通用分析数据集 → 全量素材前30秒抽帧 / 镜头切换 / OCR / 复核拼图。

用法:
  python3 pipeline.py prepare --source <xlsx>
  python3 pipeline.py media --workers 8
  python3 pipeline.py ocr
  python3 pipeline.py sheets
"""
import argparse
import concurrent.futures as cf
import hashlib
import io
import json
import math
import re
import subprocess
import tempfile
import threading
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import cv2
import imageio_ffmpeg
import numpy as np
import requests
from openpyxl import load_workbook
from PIL import Image, ImageDraw

BASE = Path(__file__).resolve().parent
DATA = BASE / 'data'
FRAMES = BASE / 'assets' / 'frames'
SHEETS = BASE / 'assets' / 'review'
FFMPEG = imageio_ffmpeg.get_ffmpeg_exe()
TIMES = [0.3, 1.5, 3, 5, 8, 12, 16, 20, 25, 29.5]
WINDOW = 30
COLS = {'运营二级行业名称': 'category', '素材链接': 'url', '消耗(元)': 'cost',
        '视频号混排分均值(有效曝光口径)': 'mix', '视频号混排eco总分均值(有效曝光口径)': 'eco',
        '3s快滑率(%)': 'skip', '平均播放时长': 'play', '浅层cvr(%)': 'cvr', 'ctr(%)': 'ctr'}
_LOCAL = threading.local()


def face_model():
    if not hasattr(_LOCAL, 'face'):
        _LOCAL.face = cv2.CascadeClassifier(cv2.data.haarcascades + 'haarcascade_frontalface_default.xml')
    return _LOCAL.face


def load(name, default=None):
    p = DATA / name
    return json.loads(p.read_text(encoding='utf-8')) if p.exists() else default


def save(name, value):
    DATA.mkdir(parents=True, exist_ok=True)
    tmp = DATA / (name + '.tmp')
    tmp.write_text(json.dumps(value, ensure_ascii=False, indent=1, allow_nan=False), encoding='utf-8')
    tmp.replace(DATA / name)


def num(v):
    try:
        f = float(v)
        return f if math.isfinite(f) else None
    except (TypeError, ValueError):
        return None


def prepare(source):
    wb = load_workbook(source, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].values)
    header = [str(h).strip() if h is not None else '' for h in rows[0]]
    miss = set(COLS) - set(header)
    if miss:
        raise SystemExit('缺少字段: %s' % miss)
    idx = {v: header.index(k) for k, v in COLS.items()}
    materials, excluded = [], []
    for rownum, row in enumerate(rows[1:], 2):
        r = {k: row[i] for k, i in idx.items()}
        if all(v in (None, '') for v in r.values()):
            continue
        url = r['url'] if isinstance(r['url'], str) else ''
        rec = {'key': f'r{rownum}', 'source_row': rownum, 'category': r['category'], 'url': url}
        for k in ['cost', 'mix', 'eco', 'skip', 'play', 'cvr', 'ctr']:
            rec[k] = num(r[k])
        if not url.startswith('http') or rec['category'] not in ('鞋靴', '箱包'):
            rec['reason'] = '整体汇总行' if r['category'] == '整体' else '无素材链接或品类不在分析范围'
            excluded.append(rec)
            continue
        m = parse_qs(urlparse(url).query).get('m', [''])[0]
        rec['id'] = m if re.fullmatch(r'[0-9a-zA-Z]{8,64}', m) else hashlib.sha256(url.encode()).hexdigest()[:32]
        materials.append(rec)
    ids = [m['id'] for m in materials]
    audit = {'source': str(source), 'sheet': wb.sheetnames[0], 'source_rows': len(rows) - 1,
             'materials': len(materials), 'excluded': excluded,
             'duplicate_ids': len(ids) - len(set(ids)),
             'missing_metrics': [m['key'] for m in materials if None in (m['eco'], m['skip'], m['play'])]}
    save('materials.json', materials)
    save('source_audit.json', audit)
    print(json.dumps({k: v for k, v in audit.items() if k != 'excluded'}, ensure_ascii=False))


def thumb(img_bgr, width=360):
    h, w = img_bgr.shape[:2]
    return cv2.resize(img_bgr, (width, int(round(h * width / w / 2)) * 2), interpolation=cv2.INTER_AREA)


def faces(img_bgr):
    g = cv2.cvtColor(img_bgr, cv2.COLOR_BGR2GRAY)
    return int(len(face_model().detectMultiScale(g, 1.1, 6, minSize=(36, 36))))


def media_one(r):
    key = r['key']
    out = FRAMES / key
    out.mkdir(parents=True, exist_ok=True)
    res = {'key': key, 'ok': False, 'kind': None, 'frames': [], 'error': None}
    url = r['url'].replace('http://', 'https://', 1)
    tmp = Path(tempfile.gettempdir()) / f'cbclip_{key}.mp4'
    try:
        with requests.get(url, stream=True, timeout=30) as resp:
            resp.raise_for_status()
            ctype = resp.headers.get('Content-Type', '')
            if ctype.startswith('image/'):
                img = Image.open(io.BytesIO(resp.content)).convert('RGB')
                arr = cv2.cvtColor(np.array(img), cv2.COLOR_RGB2BGR)
                dst = out / '00.jpg'
                cv2.imwrite(str(dst), thumb(arr), [cv2.IMWRITE_JPEG_QUALITY, 85])
                res.update(ok=True, kind='image', frames=[{'t': 0, 'path': str(dst.relative_to(BASE))}],
                           face_auto=[faces(thumb(arr))], content_type=ctype)
                return res
        res['kind'] = 'video'
        cmd = [FFMPEG, '-hide_banner', '-nostdin', '-y', '-rw_timeout', '30000000', '-i', url,
               '-t', str(WINDOW), '-map', '0:v:0', '-c', 'copy', '-an', str(tmp)]
        p = subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, timeout=420)
        log = p.stderr.decode(errors='replace')
        m = re.search(r'Duration: (\d+):(\d+):(\d+(?:\.\d+)?)', log)
        if m:
            h, mi, s = map(float, m.groups())
            res['duration'] = round(h * 3600 + mi * 60 + s, 2)
        if p.returncode != 0 or not tmp.exists() or tmp.stat().st_size < 2000:
            raise RuntimeError('下载失败: ' + log[-300:])
        cap = cv2.VideoCapture(str(tmp))
        fps = cap.get(cv2.CAP_PROP_FPS) or 25
        step = max(1, int(round(fps / 5)))
        targets = list(TIMES)
        sampled, prev_hist, cuts, diffs = [], None, [], []
        i = 0
        while True:
            ok = cap.grab()
            if not ok:
                break
            t = i / fps
            if t > WINDOW + 0.05:
                break
            need_target = targets and t >= targets[0]
            if i % step == 0 or need_target:
                ok, frame = cap.retrieve()
                if not ok:
                    break
                if i % step == 0:
                    small = cv2.resize(frame, (64, 112))
                    hsv = cv2.cvtColor(small, cv2.COLOR_BGR2HSV)
                    hist = cv2.calcHist([hsv], [0, 1], None, [16, 8], [0, 180, 0, 256])
                    cv2.normalize(hist, hist)
                    gray = cv2.cvtColor(small, cv2.COLOR_BGR2GRAY).astype(np.float32)
                    if prev_hist is not None:
                        corr = cv2.compareHist(prev_hist[0], hist, cv2.HISTCMP_CORREL)
                        diff = float(np.abs(gray - prev_hist[1]).mean())
                        diffs.append(diff)
                        if corr < 0.6 and diff > 18:
                            cuts.append(round(t, 2))
                    prev_hist = (hist, gray)
                while targets and t >= targets[0]:
                    tt = targets.pop(0)
                    dst = out / f'{len(sampled):02d}.jpg'
                    small_frame = thumb(frame)
                    cv2.imwrite(str(dst), small_frame, [cv2.IMWRITE_JPEG_QUALITY, 82])
                    sampled.append({'t': tt, 'path': str(dst.relative_to(BASE)), 'faces': faces(small_frame)})
            i += 1
        cap.release()
        clip_len = i / fps
        merged = []
        for c in cuts:
            if not merged or c - merged[-1] > 0.45:
                merged.append(c)
        res.update(frames=[{'t': f['t'], 'path': f['path']} for f in sampled],
                   face_auto=[f['faces'] for f in sampled], clip_len=round(clip_len, 2),
                   cuts15=sum(c <= 15 for c in merged), cuts30=len(merged), cut_times=merged,
                   static_ratio=round(float(np.mean([d < 3 for d in diffs])), 3) if diffs else None)
        res['ok'] = len(sampled) >= min(4, len(TIMES)) or (clip_len < 5 and len(sampled) >= 2)
        if not res['ok']:
            res['error'] = f'有效帧不足（{len(sampled)}）'
    except Exception as e:
        res['error'] = str(e)[:400]
    finally:
        tmp.unlink(missing_ok=True)
    return res


def media(workers, retry):
    mats = load('materials.json')
    done = load('media.json', {})
    todo = [m for m in mats if not done.get(m['key'], {}).get('ok') and (retry or m['key'] not in done)]
    print('待处理', len(todo), flush=True)
    with cf.ThreadPoolExecutor(workers) as pool:
        futs = {pool.submit(media_one, m): m for m in todo}
        for n, f in enumerate(cf.as_completed(futs), 1):
            r = f.result()
            done[r['key']] = r
            if n % 10 == 0 or n == len(futs):
                save('media.json', done)
            print(n, r['key'], r['ok'], r['kind'], (r['error'] or '')[:80], flush=True)
    save('media.json', done)
    print('成功', sum(v['ok'] for v in done.values()), '/', len(mats), flush=True)


SWIFT = r'''
import Foundation
import Vision
let paths = try! String(contentsOfFile: CommandLine.arguments[1]).split(separator: "\n").map(String.init)
var out: [String: String] = [:]
for p in paths {
  guard let img = NSImage(contentsOfFile: p), let cg = img.cgImage(forProposedRect: nil, context: nil, hints: nil) else { continue }
  let req = VNRecognizeTextRequest()
  req.recognitionLevel = .accurate
  req.recognitionLanguages = ["zh-Hans", "en-US"]
  req.usesLanguageCorrection = true
  try? VNImageRequestHandler(cgImage: cg).perform([req])
  let lines = (req.results ?? []).compactMap { o -> String? in
    guard let c = o.topCandidates(1).first, c.confidence > 0.4 else { return nil }
    return c.string }
  out[p] = lines.joined(separator: " / ")
}
let d = try! JSONSerialization.data(withJSONObject: out, options: [])
FileManager.default.createFile(atPath: CommandLine.arguments[2], contents: d)
'''


def ocr():
    media_ = load('media.json', {})
    done = load('ocr.json', {})
    src = BASE / 'tools' / 'ocr.swift'
    binp = BASE / 'tools' / 'ocr_bin'
    src.parent.mkdir(exist_ok=True)
    src.write_text(SWIFT.replace('import Vision', 'import Vision\nimport AppKit'))
    if not binp.exists():
        subprocess.run(['swiftc', '-O', str(src), '-o', str(binp)], check=True)
    paths = [f['path'] for m in media_.values() if m.get('ok') for f in m['frames'] if f['path'] not in done]
    print('待识别', len(paths), flush=True)
    for s in range(0, len(paths), 200):
        batch = paths[s:s + 200]
        lst = Path(tempfile.gettempdir()) / 'cb_ocr_list.txt'
        res = Path(tempfile.gettempdir()) / 'cb_ocr_out.json'
        lst.write_text('\n'.join(str(BASE / p) for p in batch))
        subprocess.run([str(binp), str(lst), str(res)], check=True, timeout=1800)
        got = json.loads(res.read_text())
        for p in batch:
            done[p] = got.get(str(BASE / p), '')
        save('ocr.json', done)
        print('OCR', min(s + 200, len(paths)), '/', len(paths), flush=True)


def sheets(per=5):
    mats = load('materials.json')
    media_ = load('media.json', {})
    text = load('ocr.json', {})
    SHEETS.mkdir(parents=True, exist_ok=True)
    (DATA / 'review_text').mkdir(exist_ok=True)
    W, TH, TW = 1090, 194, 109
    index = []
    for s in range(0, len(mats), per):
        batch = mats[s:s + per]
        no = s // per + 1
        canvas = Image.new('RGB', (W, per * (TH + 34)), '#eef1f3')
        d = ImageDraw.Draw(canvas)
        lines = []
        for j, m in enumerate(batch):
            me = media_.get(m['key'], {})
            y = j * (TH + 34)
            def f(v, fmt):
                return fmt % v if v is not None else 'NA'
            d.text((4, y + 3), f"[{j+1}] {m['key']} | ECO {f(m['eco'],'%.0f')} | skip {f(m['skip'],'%.1f')}% | play {f(m['play'],'%.1f')}s | ctr {f(m['ctr'],'%.2f')}% | cost {f(m['cost'],'%.0f')} | dur {me.get('duration','NA')}s | cuts30 {me.get('cuts30','NA')} | {me.get('kind')}", fill='black')
            for i, fr in enumerate(me.get('frames', [])[:10]):
                with Image.open(BASE / fr['path']) as im:
                    im.thumbnail((TW - 3, TH))
                    canvas.paste(im, (i * TW, y + 18))
                d.text((i * TW + 2, y + 18 + TH - 12), f"{fr['t']}s", fill='yellow')
            ocr_lines = []
            last = None
            for fr in me.get('frames', []):
                t = text.get(fr['path'], '').strip()
                if t and t != last:
                    ocr_lines.append(f"  {fr['t']}s: {t[:120]}")
                last = t
            lines.append(f"[{j+1}] {m['key']} 品类={m['category']} ECO={m['eco']} 快滑={m['skip']} 均播={m['play']} 原片时长={me.get('duration')} 媒体={me.get('kind')} 切镜30s={me.get('cuts30')} 状态={'OK' if me.get('ok') else '失败:'+str(me.get('error'))[:60]}\n" + ('\n'.join(ocr_lines) or '  (无可识别字幕)'))
        path = SHEETS / f'{no:03d}.jpg'
        canvas.save(path, quality=86)
        (DATA / 'review_text' / f'{no:03d}.txt').write_text('\n'.join(lines), encoding='utf-8')
        index.append({'sheet': str(path.relative_to(BASE)), 'text': f'data/review_text/{no:03d}.txt', 'keys': [m['key'] for m in batch]})
    save('review_index.json', index)
    print('拼图', len(index))


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('action', choices=['prepare', 'media', 'ocr', 'sheets'])
    ap.add_argument('--source', type=Path)
    ap.add_argument('--workers', type=int, default=8)
    ap.add_argument('--retry', action='store_true')
    a = ap.parse_args()
    if a.action == 'prepare':
        prepare(a.source)
    elif a.action == 'media':
        media(a.workers, a.retry)
    elif a.action == 'ocr':
        ocr()
    else:
        sheets()
